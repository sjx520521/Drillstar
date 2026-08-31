# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "DrillStar钻井数据采集分析平台V2.0-重写版.docx"
TARGET_DIR = Path(r"E:\数据分析软件DrillStar\DrillStar钻井数据采集分析平台V1.0\DrillStar钻井数据采集分析平台V1.0(软著)")


def set_run_font(run, size=10.5, bold=False, font="宋体"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def set_paragraph_format(paragraph, first_line=True, line_spacing=1.5):
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(6)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)


def add_paragraph(doc, text="", size=10.5, bold=False, align=None, first_line=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, first_line=first_line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(8)
        size = 16
    elif level == 2:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(6)
        size = 14
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        size = 12
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, font="宋体")
    return p


def add_image_placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=bold)


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for style_name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = True


def add_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tbl_pr.append(borders)


def build_doc():
    doc = Document()
    set_doc_defaults(doc)

    add_paragraph(doc, "DrillStar钻井数据采集分析平台V2.0", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(doc, "使用手册", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(doc, "", first_line=False)

    table = doc.add_table(rows=6, cols=2)
    add_table_borders(table)
    rows = [
        ("版本", "2.0"),
        ("文档状态", "编辑"),
        ("作者", ""),
        ("负责人", ""),
        ("创建日期", "2026年6月15日"),
        ("更新日期", "2026年6月15日"),
    ]
    for row, (key, value) in zip(table.rows, rows):
        set_cell_text(row.cells[0], key, bold=True)
        set_cell_text(row.cells[1], value)

    add_page_break(doc)

    add_heading(doc, "目录", 1)
    toc_lines = [
        "1. 简介",
        "1.1 编写目的",
        "2. 软件概述",
        "3. 运行及开发环境",
        "3.1 运行环境",
        "3.2 部署环境",
        "3.3 硬件与接口环境",
        "3.4 开发环境",
        "4. 使用说明",
        "4.1 首页界面",
        "4.2 模拟实验模块",
        "4.3 数据管理模块",
        "4.4 数据分析模块",
        "4.4.1 数据导入与数据整理",
        "4.4.2 数据预处理与信号分析",
        "4.4.3 工程专项分析",
        "4.4.4 机器学习模型分析",
        "4.4.5 模型保存、加载与预测",
        "4.4.6 结果显示与保存",
        "4.5 数据分析链路",
        "4.6 使用注意事项",
    ]
    for line in toc_lines:
        add_paragraph(doc, line, first_line=False)

    add_page_break(doc)

    add_heading(doc, "1. 简介", 1)
    add_paragraph(doc, "DrillStar钻井数据采集分析平台V2.0是面向钻井实验与工程数据处理场景研发的桌面应用软件。软件以Python为主要开发语言，以PyQt5为图形界面框架，围绕钻压、扭矩、转速、振动、压力、温度、进尺、泵压、流量等钻井关键参数，提供实时采集、数据管理、节点式数据分析、工程专项计算、机器学习模型分析、结果可视化和数据导出等功能。")
    add_paragraph(doc, "本软件采用节点式流程组织方式，将数据导入、数据表整理、数据清洗、滤波、时域分析、频域分析、时频分析、小波分析、重采样、工程计算、模型训练、模型保存、模型加载、模型预测和结果输出等功能封装为可连接的分析节点。用户可根据不同实验任务和工程分析目标，自主搭建完整的数据分析链路，实现从原始数据到分析结果的连续处理。")
    add_paragraph(doc, "软件既可用于钻井实验现场的实时监测，也可用于历史数据的离线整理、专题分析和模型预测。通过文件数据、数据库数据和实时采集数据的统一接入，软件能够为实验复盘、工况分析、异常识别、风险研判和模型验证提供一体化的数据处理环境。")

    add_heading(doc, "1.1 编写目的", 2)
    add_paragraph(doc, "本文档用于说明DrillStar钻井数据采集分析平台V2.0的软件组成、运行环境、功能模块、典型数据分析链路及使用注意事项。文档在介绍软件基础功能和工程功能的同时，重点说明机器学习模型分析、模型保存加载、预测节点以及多类型数据分析链路的实现方式。")
    add_paragraph(doc, "本文档可作为软件使用说明、功能说明和软件著作权登记材料的组成部分，供软件使用者、维护人员和相关评审人员了解软件的功能范围、操作方式和技术特点。")

    add_heading(doc, "2. 软件概述", 1)
    add_paragraph(doc, "DrillStar钻井数据采集分析平台V2.0采用分层、模块化设计。系统整体围绕“数据接入、数据管理、数据分析、模型应用、结果输出”的业务过程展开，形成从实验数据采集、历史数据调用、数据预处理、工程特征计算到模型训练和预测输出的完整工作流程。")
    add_paragraph(doc, "软件首页提供模拟实验、数据管理和数据分析三个主入口。模拟实验模块主要完成实时数据接入和曲线监测；数据管理模块主要完成数据库连接、数据上传下载和历史数据维护；数据分析模块主要完成离线数据导入、节点式流程搭建、信号处理、工程专项分析、机器学习模型训练、模型复用预测和结果保存。")
    add_paragraph(doc, "在数据分析模块中，软件通过节点编辑器组织分析流程。每个功能节点具有明确的输入和输出，用户通过连线确定数据流向。该方式使软件不仅能够完成单项功能操作，还能够将多个功能组合为完整链路，适应不同类型的钻井数据分析任务。")
    add_image_placeholder(doc, "【图2.1 系统总体框架图】")

    add_heading(doc, "3. 运行及开发环境", 1)
    add_heading(doc, "3.1 运行环境", 2)
    add_paragraph(doc, "本软件运行于Windows 10或Windows 11 64位操作系统，采用桌面应用程序形式部署。软件打包后可直接在目标计算机上运行。")
    add_heading(doc, "3.2 部署环境", 2)
    add_paragraph(doc, "在发布部署场景中，软件可打包为Windows可执行程序。若需要使用数据库功能，应保证目标计算机能够访问MySQL数据库；若需要使用NI采集功能，应安装相应NI驱动和运行环境。")
    add_heading(doc, "3.3 硬件与接口环境", 2)
    add_paragraph(doc, "在实时采集场景中，软件支持TCP网络接口和NI DAQmx采集接口，可结合实验台架、传感器采集设备和数据库系统完成数据接入。")
    add_heading(doc, "3.4 开发环境", 2)
    add_paragraph(doc, "本软件使用Python作为主要开发语言，使用PyCharm作为主要开发工具，界面部分通过Qt Designer设计，并基于PyQt5完成界面集成与业务逻辑开发。数据处理主要依赖Pandas、NumPy和SciPy，图形显示主要依赖PyQtGraph和Matplotlib，机器学习模型主要依赖scikit-learn，模型保存与加载使用joblib。")
    add_paragraph(doc, "软件开发和运行中使用的主要包包括：PyQt5、QtPy、pyqtgraph、matplotlib、numpy、pandas、scipy、scikit-learn、joblib、openpyxl、PyWavelets、nidaqmx、PyMySQL、cryptography、python-dateutil和PyInstaller等。其中LSTM和Informer等深度学习模型需要PyTorch环境支持。")

    add_heading(doc, "4. 使用说明", 1)
    add_heading(doc, "4.1 首页界面", 2)
    add_paragraph(doc, "软件首页提供模拟实验、数据管理和数据分析三个入口。用户可根据当前任务选择进入对应模块。若需要进行实时采集与曲线监测，可进入模拟实验模块；若需要进行数据库数据维护，可进入数据管理模块；若需要对历史数据进行处理、分析、建模或预测，可进入数据分析模块。")
    add_image_placeholder(doc, "【图4.1 首页界面】")

    add_heading(doc, "4.2 模拟实验模块", 2)
    add_paragraph(doc, "模拟实验模块用于钻井实验参数的实时采集、接入与动态显示。该模块支持TCP网络连接和NI DAQmx采集两种数据接入方式，可对钻压、扭矩、转速、振动、压力、温度、进尺、泵压、流量等参数进行实时监测。用户可在参数树中选择需要显示的参数，并通过曲线窗口观察实验过程中的动态变化。")
    add_paragraph(doc, "该模块支持多个参数窗口同时打开，并可对曲线颜色、比例系数和线宽等显示属性进行调整。通过实时曲线显示和窗口管理功能，用户能够对实验过程中的关键参数变化进行并行观察，为后续数据分析和工程判断提供基础依据。")
    add_image_placeholder(doc, "【图4.2 模拟实验模块界面】")

    add_heading(doc, "4.3 数据管理模块", 2)
    add_paragraph(doc, "数据管理模块用于实验数据和历史数据的集中存储、查询与维护。该模块支持MySQL数据库连接，能够读取数据库表列表，查看表信息，完成本地数据上传、数据库数据下载以及数据表维护等操作。")
    add_paragraph(doc, "通过数据管理模块，软件能够实现本地文件、数据库和分析流程之间的数据流转。用户既可以将实验数据保存到数据库中，也可以将数据库中的历史数据调入分析模块进行复用分析。")
    add_image_placeholder(doc, "【图4.3 数据管理模块界面】")

    add_heading(doc, "4.4 数据分析模块", 2)
    add_paragraph(doc, "数据分析模块是本软件的核心模块。该模块采用节点式流程组织方式，将数据导入、数据整理、预处理、工程分析、机器学习模型、结果显示和结果保存等功能封装为不同节点。用户可根据分析目标创建节点并连接数据流向，构建适合当前任务的数据分析链路。")
    add_image_placeholder(doc, "【图4.4 数据分析模块主界面】")

    add_heading(doc, "4.4.1 数据导入与数据整理", 3)
    add_paragraph(doc, "数据导入功能用于将不同来源的数据接入分析流程。系统支持CSV、Excel和SQL三种数据导入方式，能够满足本地文件数据和数据库历史数据的导入需求。导入后的数据可作为后续数据清洗、信号处理、工程计算和模型训练的输入。")
    add_image_placeholder(doc, "【图4.5 数据导入界面】")
    add_paragraph(doc, "数据表功能用于以表格形式查看和整理导入后的数据。用户可选择参与分析的列，设置行范围，预览当前数据内容，并对单元格、列标题、行和列进行必要调整。数据信息功能用于查看数据的基本属性和结构信息，便于用户在正式分析前了解数据组成情况。")
    add_image_placeholder(doc, "【图4.6 数据信息与数据表界面】")

    add_heading(doc, "4.4.2 数据预处理与信号分析", 3)
    add_paragraph(doc, "数据预处理功能用于改善原始数据质量并生成适合后续分析的数据。数据清洗节点可处理缺失值和异常值，滤波节点可对数值信号进行平滑处理，重采样与时间对齐节点可统一不同数据源的采样间隔和时间基准。")
    add_image_placeholder(doc, "【图4.7 数据清洗界面】")
    add_image_placeholder(doc, "【图4.8 数据滤波界面】")
    add_paragraph(doc, "时域分析节点用于显示多参数时间序列曲线，并支持用户选择分析区间，计算最大值、最小值、平均值、方差、标准差、均方根等时域统计特征。频域分析节点用于对选定数据执行FFT分析，时频分析节点用于执行STFT分析，小波分析节点用于小波分解和小波去噪处理。这些功能可用于振动信号、扭矩波动、转速变化等钻井时序信号的多角度分析。")
    add_image_placeholder(doc, "【图4.9 时域分析界面】")
    add_image_placeholder(doc, "【图4.10 频域分析界面】")
    add_image_placeholder(doc, "【图4.11 时频分析界面】")
    add_image_placeholder(doc, "【图4.12 小波分析界面】")
    add_paragraph(doc, "重采样与时间对齐功能可对单源数据进行升采样或降采样，也可对两组数据按照时间轴进行对齐。该功能适用于多传感器数据融合、不同采样频率数据统一和模型训练前的数据序列整理。")
    add_image_placeholder(doc, "【图4.13 重采样与时间对齐界面】")

    add_heading(doc, "4.4.3 工程专项分析", 3)
    add_paragraph(doc, "工程专项分析功能面向钻井实验与工程数据处理场景，主要包括工程值转换、时间转换、时深对标、机械比能MSE计算、粘滑分析和卡钻分析。该类功能能够将基础参数转化为具有工程含义的分析结果，为实验复盘和工况研判提供依据。")
    add_paragraph(doc, "机械比能MSE计算功能根据钻压、扭矩、钻速、钻头直径等参数计算机械比能，并生成MSE结果列和曲线显示结果。该功能可用于钻进效率分析、钻头状态判断和地层变化辅助识别。")
    add_image_placeholder(doc, "【图4.14 MSE计算界面与结果图】")
    add_paragraph(doc, "粘滑分析功能用于识别钻井过程中的粘滑振动状态。卡钻分析功能根据钻压、扭矩、转速等参数变化计算卡钻风险值、风险趋势和风险等级。时深对标功能用于将时间序列数据与深度数据建立对应关系，便于按井深维度观察参数变化。")
    add_image_placeholder(doc, "【图4.15 粘滑分析界面与结果图】")
    add_image_placeholder(doc, "【图4.16 卡钻分析界面与结果图】")
    add_image_placeholder(doc, "【图4.17 时深对标界面与结果图】")

    add_heading(doc, "4.4.4 机器学习模型分析", 3)
    add_paragraph(doc, "机器学习模型分析功能用于在数据预处理和特征构建的基础上，对钻井实验数据开展降维、聚类、分类和回归分析。用户可将清洗后的数据、时域特征、MSE结果、小波处理结果或人工整理后的数据表接入模型节点，选择特征列和目标列，设置模型参数，并在后台完成训练或计算。")
    add_paragraph(doc, "PCA主成分分析节点用于对多维数值特征进行降维处理。该节点可将多参数数据转换为主成分列，用于高维数据可视化、后续聚类分析或模型训练前的数据压缩。")
    add_image_placeholder(doc, "【图4.18 PCA模型运行界面与结果图】")
    add_paragraph(doc, "KMeans聚类节点用于无标签数据的状态分群。用户可将多参数特征输入该节点，设置聚类数量和迭代参数，系统会输出聚类标签和距离信息，可用于工况初步划分、异常样本发现和人工标注前的数据探索。")
    add_image_placeholder(doc, "【图4.19 KMeans模型运行界面与结果图】")
    add_paragraph(doc, "SVM支持向量机节点用于有标签分类任务。该节点适合用于正常与异常状态识别、工况类别判断、粘滑或卡钻风险类别识别等场景。系统支持特征标准化、核函数选择、参数配置和分类结果输出。")
    add_image_placeholder(doc, "【图4.20 SVM模型运行界面与结果图】")
    add_paragraph(doc, "线性回归节点用于连续值预测任务，可用于钻速、扭矩、MSE等连续参数的回归分析。节点支持普通最小二乘、Ridge、Lasso和ElasticNet等模型形式，并输出预测值、残差和评价指标。")
    add_image_placeholder(doc, "【图4.21 线性回归模型运行界面与结果图】")
    add_paragraph(doc, "逻辑回归节点用于分类任务，适合对可解释性要求较高的工况分类分析。该节点可根据用户选择的特征列和目标列完成模型训练，并输出预测类别和分类报告。")
    add_image_placeholder(doc, "【图4.22 逻辑回归模型运行界面与结果图】")
    add_paragraph(doc, "朴素贝叶斯节点用于快速分类分析，支持Gaussian、Multinomial、Bernoulli和Complement等形式。该节点适合在样本规模较小或需要快速建立分类基线模型时使用。")
    add_image_placeholder(doc, "【图4.23 朴素贝叶斯模型运行界面与结果图】")
    add_paragraph(doc, "神经网络节点用于处理较复杂的非线性分类或回归关系。用户可通过该节点对多参数钻井数据进行非线性建模，并将预测结果传递给后续数据表、图表或保存节点。")
    add_image_placeholder(doc, "【图4.24 神经网络模型运行界面与结果图】")
    add_paragraph(doc, "LSTM和Informer节点用于时间序列建模。该类节点适合处理具有时间依赖关系的参数序列，可用于连续钻井参数预测或较长序列趋势分析。使用该类模型时，运行环境需要具备相应深度学习库支持。")
    add_image_placeholder(doc, "【图4.25 LSTM模型运行界面与结果图】")
    add_image_placeholder(doc, "【图4.26 Informer模型运行界面与结果图】")

    add_heading(doc, "4.4.5 模型保存、加载与预测", 3)
    add_paragraph(doc, "为使机器学习模型能够在不同任务和不同数据批次中复用，软件提供模型包对象、模型保存节点、模型加载节点和模型预测节点。模型包对象用于保存模型本身、特征列、目标列、参数配置、预处理器、标签编码器和创建时间等信息，从而保证模型在保存、加载和预测时保持输入字段和处理方式一致。")
    add_paragraph(doc, "SaveModel节点用于保存已训练模型。用户将训练完成的模型节点连接到SaveModel节点后，可在窗口中加载上游模型包，并将其保存为本地模型文件。该功能使训练结果能够持久化保存，便于后续分析任务复用。")
    add_image_placeholder(doc, "【图4.27 SaveModel保存模型节点界面】")
    add_paragraph(doc, "LoadModel节点用于加载已保存的模型包。用户可打开历史模型文件，系统会校验模型包有效性并显示模型摘要信息。加载后的模型包可传递给Predict节点，用于新数据预测。")
    add_image_placeholder(doc, "【图4.28 LoadModel加载模型节点界面】")
    add_paragraph(doc, "Predict节点用于使用模型包对新数据进行预测。该节点包含两个输入端，其中输入1连接模型包，输入2连接待预测数据。预测完成后，系统会在原始数据表中追加预测列、置信度列、聚类标签列或主成分列，并将结果传递到下游节点。")
    add_image_placeholder(doc, "【图4.29 Predict模型预测节点界面与预测结果图】")

    add_heading(doc, "4.4.6 结果显示与保存", 3)
    add_paragraph(doc, "结果显示功能用于对分析结果进行图形化展示。软件支持折线图、散点图和柱状图，可用于查看时序趋势、参数相关性、分类或聚类结果分布等内容。图表节点可接收数据处理节点、工程分析节点或预测节点输出的数据表。")
    add_image_placeholder(doc, "【图4.30 折线图、散点图和柱状图显示界面】")
    add_paragraph(doc, "结果保存功能用于将处理结果、工程分析结果或预测结果导出为CSV或Excel文件。通过保存节点，用户可将软件内部生成的数据表输出到本地文件，用于报告编写、数据归档或后续分析。")
    add_image_placeholder(doc, "【图4.31 结果保存界面】")

    add_heading(doc, "4.5 数据分析链路", 2)
    add_paragraph(doc, "数据分析链路是本软件数据分析模块的重要特点。用户可将不同功能节点按照分析目标连接起来，使数据在节点之间逐步流动，最终形成从原始数据接入到结果输出的完整处理流程。以下为软件可实现的典型数据分析链路。")

    add_heading(doc, "4.5.1 数据清洗与可视化链路", 3)
    add_paragraph(doc, "该链路通常由CSV导入、Excel导入或SQL导入节点开始，数据进入数据清洗节点后完成缺失值和异常值处理，再进入滤波节点进行平滑处理，最后通过折线图、散点图或柱状图进行显示，并通过保存节点输出处理结果。该链路适用于传感器原始数据整理、实验过程曲线查看和分析前数据质量改善。")
    add_paragraph(doc, "链路形式：CSV/Excel/SQL导入 → 数据清洗 → 滤波 → 折线图/散点图/柱状图 → 保存数据。")
    add_image_placeholder(doc, "【图4.32 数据清洗与可视化链路图及运行结果图】")

    add_heading(doc, "4.5.2 时域与信号处理链路", 3)
    add_paragraph(doc, "该链路面向钻井时序信号分析。用户可将清洗后的数据送入时域分析、频域分析、时频分析或小波分析节点，从不同角度观察参数变化和信号特征。时域分析可输出统计特征，频域分析可识别频率分布，时频分析可观察频率随时间变化，小波分析可用于非平稳信号去噪和多尺度分析。")
    add_paragraph(doc, "链路形式：数据导入 → 数据清洗 → 时域分析/频域分析/时频分析/小波分析 → 图表显示或结果保存。")
    add_image_placeholder(doc, "【图4.33 时域与信号处理链路图及运行结果图】")

    add_heading(doc, "4.5.3 工程指标分析链路", 3)
    add_paragraph(doc, "该链路面向钻井工程指标计算和风险分析。数据经过导入、清洗和必要滤波后，可进入MSE计算、粘滑分析、卡钻分析或时深对标节点。分析结果可继续进行图形显示，也可作为后续机器学习模型的输入特征。")
    add_paragraph(doc, "链路形式：数据导入 → 数据清洗 → 滤波 → MSE计算/粘滑分析/卡钻分析/时深对标 → 图表显示 → 保存数据。")
    add_image_placeholder(doc, "【图4.34 工程指标分析链路图及运行结果图】")

    add_heading(doc, "4.5.4 工况聚类分析链路", 3)
    add_paragraph(doc, "该链路用于无标签数据的工况划分。用户可先通过清洗、滤波、MSE计算、时域特征或小波分析获得特征数据，再将特征数据输入KMeans节点进行聚类。聚类结果可通过散点图观察不同状态分布，并可保存为后续人工标注或监督学习训练数据。")
    add_paragraph(doc, "链路形式：数据导入 → 数据清洗 → 特征构建 → KMeans聚类 → 散点图/数据表 → 保存数据。")
    add_image_placeholder(doc, "【图4.35 工况聚类分析链路图及运行结果图】")

    add_heading(doc, "4.5.5 监督学习训练链路", 3)
    add_paragraph(doc, "该链路用于基于历史数据训练分类或回归模型。用户可将历史数据导入后进行清洗、滤波和特征构建，再通过数据表节点整理特征列和目标列，随后接入SVM、逻辑回归、线性回归、朴素贝叶斯或神经网络节点进行训练。训练完成后，可将模型连接到SaveModel节点保存为模型文件。")
    add_paragraph(doc, "链路形式：历史数据导入 → 数据清洗 → 特征构建 → 模型训练 → SaveModel保存模型。")
    add_image_placeholder(doc, "【图4.36 监督学习训练链路图及模型运行结果图】")

    add_heading(doc, "4.5.6 模型复用预测链路", 3)
    add_paragraph(doc, "该链路用于将已保存模型应用于新数据。用户通过LoadModel节点加载历史模型包，同时将新数据通过导入、清洗和特征构建节点处理为与训练时一致的数据形式，然后将模型包和新数据分别接入Predict节点。预测结果可传递到数据表、图表或保存节点。")
    add_paragraph(doc, "链路形式：LoadModel → Predict输入1；新数据导入 → 数据清洗/特征构建 → Predict输入2；Predict → 图表显示/保存数据。")
    add_image_placeholder(doc, "【图4.37 模型复用预测链路图及预测结果图】")

    add_heading(doc, "4.5.7 综合数据分析链路", 3)
    add_paragraph(doc, "在较完整的分析任务中，用户可将多个节点组合为综合链路。例如，先从数据库或本地文件导入数据，经数据表整理、数据清洗、重采样与时间对齐、滤波和MSE计算后，进一步进行时域特征提取、PCA降维、KMeans聚类或SVM分类模型训练。训练完成的模型可保存为模型包，并在后续新数据中加载使用。该链路体现了软件从数据接入、处理、工程计算、模型分析到预测输出的完整能力。")
    add_paragraph(doc, "链路形式：数据导入 → 数据表整理 → 数据清洗 → 重采样与时间对齐 → 滤波 → MSE计算 → 时域特征分析 → PCA/KMeans/SVM等模型分析 → SaveModel/LoadModel/Predict → 图表显示 → 保存数据。")
    add_image_placeholder(doc, "【图4.38 综合数据分析链路图及运行结果图】")

    add_heading(doc, "4.6 使用注意事项", 2)
    add_paragraph(doc, "使用节点式分析流程时，应保证上游节点输出内容与下游节点输入要求一致。图表节点通常需要数据表输入，模型训练节点需要包含特征列和目标列的数据表，Predict节点的输入1需要模型包，输入2需要待预测数据表。")
    add_paragraph(doc, "进行机器学习模型训练前，应确认数据列类型、特征列、目标列和单位含义正确。若模型训练时使用了特定特征列，新数据预测时也应包含相同特征列，否则预测节点会提示字段缺失。")
    add_paragraph(doc, "对于数据量较大的文件，建议先通过数据表、重采样或特征提取节点减少数据规模，再进行复杂图形显示或模型训练。对于LSTM和Informer等深度学习模型，应确认运行环境已安装相应深度学习库。")

    doc.save(OUT)
    target = TARGET_DIR / "DrillStar钻井数据采集分析平台V2.0-重写版.docx"
    if TARGET_DIR.exists():
        doc.save(target)
    print(OUT)
    if TARGET_DIR.exists():
        print(target)


if __name__ == "__main__":
    build_doc()
