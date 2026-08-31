import logging
import os

import numpy as np
import pandas as pd
from PyQt5.QtCore import QAbstractTableModel, QThread, QTimer, Qt, pyqtSignal, pyqtSlot
from PyQt5.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QHeaderView,
    QMessageBox,
    QWidget,
)

try:
    from openpyxl import load_workbook
except ImportError as exc:
    load_workbook = None
    OPENPYXL_IMPORT_ERROR = exc
else:
    OPENPYXL_IMPORT_ERROR = None

try:
    import xlrd  # noqa: F401
except ImportError as exc:
    xlrd = None
    XLRD_IMPORT_ERROR = exc
else:
    XLRD_IMPORT_ERROR = None

from dataanalysis.buttonNode import ButtonNode
from dataanalysis.ui.csvImport_ui import Ui_CSVImport


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    handlers=[logging.StreamHandler(), logging.FileHandler("excel_import.log")],
)
logger = logging.getLogger(__name__)

WINDOW_TITLE = "DataAnalysis [Excel Import]"
FILE_FILTER = "Excel Files (*.xlsx *.xls *.xlsm);;All Files (*.*)"
DEFAULT_FILE_PATH = os.path.expanduser("~")
MAX_DISPLAY_ROWS = 1000


class ExcelDependencyError(ImportError):
    """Raised when the required Excel parser dependency is unavailable."""


class ExcelLoaderThread(QThread):
    progress_update = pyqtSignal(int)
    data_loaded = pyqtSignal(pd.DataFrame)
    error_occurred = pyqtSignal(str)

    def __init__(self, file_path: str):
        super().__init__()
        self.file_path = file_path
        self._is_running = True

    def stop(self):
        self._is_running = False

    def _is_continuous_increment(self, series: pd.Series) -> bool:
        series = series.dropna()
        if len(series) == 0:
            return False

        if not pd.api.types.is_integer_dtype(series.dtype):
            if not pd.api.types.is_float_dtype(series.dtype):
                return False
            if not np.all(series == series.astype(int)):
                return False

        arr = series.astype(int).to_numpy()
        return np.array_equal(arr, np.arange(arr[0], arr[0] + len(arr)))

    def _remove_auto_increment_column(self, df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df

        if df.columns[0] == 0 or str(df.columns[0]).isdigit():
            return df.drop(columns=[df.columns[0]])

        for col in df.columns:
            if self._is_continuous_increment(df[col]):
                return df.drop(columns=[col])
        return df

    def _optimize_dtypes(self, df: pd.DataFrame) -> pd.DataFrame:
        for col in df.columns:
            if df[col].dtype == "float64":
                clean = df[col].dropna()
                if len(clean) > 0 and np.all(clean == clean.astype(int)):
                    if df[col].isnull().any():
                        try:
                            df[col] = df[col].astype("Int64")
                        except Exception:
                            pass
                    else:
                        df[col] = df[col].astype("int32")
            elif df[col].dtype == "int64":
                df[col] = pd.to_numeric(df[col], downcast="integer")
        return df

    def _read_xlsx_fast(self) -> pd.DataFrame:
        if load_workbook is None:
            raise ExcelDependencyError("读取 .xlsx/.xlsm 文件需要 openpyxl，请先安装 openpyxl") from OPENPYXL_IMPORT_ERROR

        workbook = load_workbook(
            filename=self.file_path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                return pd.DataFrame()

            columns = []
            for index, value in enumerate(header):
                text = "" if value is None else str(value).strip()
                columns.append(text or f"Column{index + 1}")

            data = list(rows)
            return pd.DataFrame(data, columns=columns)
        finally:
            workbook.close()

    def _read_excel(self) -> pd.DataFrame:
        file_suffix = os.path.splitext(self.file_path)[-1].lower()
        if file_suffix in {".xlsx", ".xlsm"}:
            return self._read_xlsx_fast()
        if file_suffix == ".xls":
            if xlrd is None:
                raise ExcelDependencyError("读取 .xls 文件需要 xlrd，请先安装 xlrd==2.0.1") from XLRD_IMPORT_ERROR
            return pd.read_excel(self.file_path, engine="xlrd")
        raise ValueError("不支持的文件格式，请选择 .xlsx、.xls 或 .xlsm 文件")

    def run(self):
        try:
            self.progress_update.emit(10)
            if not self._is_running:
                return

            self.progress_update.emit(25)
            df = self._read_excel()

            if not self._is_running:
                return

            self.progress_update.emit(75)
            df = self._remove_auto_increment_column(df)
            df = self._optimize_dtypes(df)

            self.progress_update.emit(100)
            self.data_loaded.emit(df)

        except pd.errors.EmptyDataError:
            self.error_occurred.emit("所选 Excel 文件为空")
        except FileNotFoundError:
            self.error_occurred.emit("文件不存在或已被删除")
        except PermissionError:
            self.error_occurred.emit("文件正在被占用，无法读取")
        except ExcelDependencyError as e:
            logger.exception("Excel parser dependency is unavailable for %s", self.file_path)
            self.error_occurred.emit(str(e))
        except ImportError:
            logger.exception("Excel parser import failed for %s", self.file_path)
            self.error_occurred.emit("缺少 Excel 解析依赖，请检查 openpyxl 或 xlrd")
        except Exception as e:
            logger.exception("Unexpected Excel import failure for %s", self.file_path)
            self.error_occurred.emit(f"读取 Excel 文件失败：{str(e)}")


class OptimizedPandasModel(QAbstractTableModel):
    def __init__(self, data: pd.DataFrame, parent=None):
        super().__init__(parent)
        self._data = data.head(MAX_DISPLAY_ROWS) if len(data) > MAX_DISPLAY_ROWS else data
        self._cols = self._data.columns.tolist()
        self._dtypes = self._data.dtypes
        self._data_cache = self._data.values.tolist()

    def rowCount(self, parent=None) -> int:
        return len(self._data_cache)

    def columnCount(self, parent=None) -> int:
        return len(self._cols)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None

        row = index.row()
        col = index.column()
        value = self._data_cache[row][col]

        if role == Qt.ItemDataRole.DisplayRole:
            if pd.isna(value):
                return ""
            if pd.api.types.is_numeric_dtype(self._dtypes.iloc[col]):
                return f"{value:.4f}" if isinstance(value, float) else str(value)
            if pd.api.types.is_datetime64_any_dtype(self._dtypes.iloc[col]):
                return value.strftime("%Y-%m-%d %H:%M:%S")
            return str(value)
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role != Qt.ItemDataRole.DisplayRole:
            return None
        if orientation == Qt.Horizontal:
            return self._cols[section] if 0 <= section < len(self._cols) else None
        return str(section + 1)

    def flags(self, index):
        return Qt.ItemFlag.ItemIsSelectable | Qt.ItemFlag.ItemIsEnabled

    def update_data(self, new_data: pd.DataFrame):
        self.beginResetModel()
        self._data = new_data.head(MAX_DISPLAY_ROWS) if len(new_data) > MAX_DISPLAY_ROWS else new_data
        self._cols = self._data.columns.tolist()
        self._dtypes = self._data.dtypes
        self._data_cache = self._data.values.tolist()
        self.endResetModel()


class DataTransferThread(QThread):
    transfer_done = pyqtSignal(bool, str)
    progress_update = pyqtSignal(int)

    def __init__(self, node: ButtonNode, df: pd.DataFrame):
        super().__init__()
        self.node = node
        self.df = df

    def run(self):
        try:
            self.progress_update.emit(20)
            if self.node.LoadData(self.df):
                self.progress_update.emit(80)
                self.node.eval()
                self.node.evalChildren()
                self.progress_update.emit(100)
                self.transfer_done.emit(True, "数据已成功传递到节点")
            else:
                self.transfer_done.emit(False, "节点 LoadData 返回失败")
        except Exception as e:
            self.transfer_done.emit(False, f"数据传递失败：{str(e)}")


class Win_ExcelImport(QWidget):
    def __init__(self, node: ButtonNode):
        super().__init__()
        self.ui = Ui_CSVImport()
        self.ui.setupUi(self)
        self._init_ui()
        self._init_variables(node)
        self._bind_signals()

        logger.info(f"Excel 导入窗口初始化，关联节点：{node}")

    def _init_ui(self):
        self.setWindowTitle(WINDOW_TITLE)
        self.setWindowModality(Qt.WindowModality.ApplicationModal)
        self.resize(1000, 600)

        self.ui.btn_OpenFile.setText("Load")
        self.ui.btn_Transfer.setText("Transfer")

        self.ui.tV_DataView.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.ui.tV_DataView.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.ui.tV_DataView.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.ui.tV_DataView.horizontalHeader().setStretchLastSection(True)
        self.ui.tV_DataView.setVerticalScrollMode(QAbstractItemView.ScrollPerItem)
        self.ui.tV_DataView.setHorizontalScrollMode(QAbstractItemView.ScrollPerItem)
        self.ui.tV_DataView.setSortingEnabled(False)

        self.ui.progressBar.setRange(0, 100)
        self.ui.progressBar.setValue(0)
        self.ui.btn_delete.setEnabled(False)
        self.ui.btn_Transfer.setEnabled(False)

    def _init_variables(self, node: ButtonNode):
        self.node = node
        self.df = None
        self.model = None
        self.loader_thread = None
        self.transfer_thread = None

    def _bind_signals(self):
        self.ui.btn_OpenFile.clicked.connect(self.on_choose_file)
        self.ui.btn_delete.clicked.connect(self.on_delete_column)
        self.ui.btn_Transfer.clicked.connect(self.on_load_data)

    @pyqtSlot()
    def on_choose_file(self):
        self.node.markInvalid()
        self.node.markDescendantsInvalid()

        filepath, _ = QFileDialog.getOpenFileName(
            self,
            "选择 Excel 文件",
            DEFAULT_FILE_PATH,
            FILE_FILTER,
        )
        if not filepath:
            return

        self.ui.cB_FilePath.clear()
        self.ui.cB_count.clear()
        self.ui.progressBar.setValue(0)
        self.ui.btn_delete.setEnabled(False)
        self.ui.btn_Transfer.setEnabled(False)
        self.ui.cB_FilePath.addItem(filepath)

        if self.loader_thread and self.loader_thread.isRunning():
            self.loader_thread.stop()
            self.loader_thread.wait()

        self.loader_thread = ExcelLoaderThread(filepath)
        self.loader_thread.progress_update.connect(self.ui.progressBar.setValue)
        self.loader_thread.data_loaded.connect(self.on_data_loaded)
        self.loader_thread.error_occurred.connect(self.on_load_error)
        self.loader_thread.start()

    @pyqtSlot(pd.DataFrame)
    def on_data_loaded(self, df: pd.DataFrame):
        if df.empty:
            QMessageBox.warning(self, "警告", "Excel 文件中没有可用数据")
            return

        self.df = df
        self.model = OptimizedPandasModel(df)
        self.ui.tV_DataView.setModel(self.model)

        QTimer.singleShot(300, lambda: self.ui.tV_DataView.setSortingEnabled(True))

        self.ui.cB_count.clear()
        for i in range(self.model.columnCount()):
            self.ui.cB_count.addItem(str(i + 1))

        has_columns = self.model.columnCount() > 0
        self.ui.btn_delete.setEnabled(has_columns)
        self.ui.btn_Transfer.setEnabled(has_columns)

        logger.info(f"Excel 文件加载完成：{df.shape}")
        QMessageBox.information(
            self,
            "成功",
            f"已加载 {len(df)} 行 × {len(df.columns)} 列数据，表格展示前 {min(len(df), MAX_DISPLAY_ROWS)} 行",
        )

    @pyqtSlot(str)
    def on_load_error(self, error_msg: str):
        logger.error(error_msg)
        QMessageBox.critical(self, "错误", error_msg)
        self.ui.progressBar.setValue(0)

    @pyqtSlot()
    def on_delete_column(self):
        if self.df is None or self.model is None:
            return

        try:
            col_index = int(self.ui.cB_count.currentText()) - 1
            if col_index < 0 or col_index >= len(self.df.columns):
                raise IndexError("列索引超出范围")

            self.df = self.df.drop(self.df.columns[col_index], axis=1)
            self.model.update_data(self.df)

            self.ui.cB_count.clear()
            for i in range(self.model.columnCount()):
                self.ui.cB_count.addItem(str(i + 1))

            has_columns = self.model.columnCount() > 0
            self.ui.btn_delete.setEnabled(has_columns)
            self.ui.btn_Transfer.setEnabled(has_columns)

        except Exception as e:
            QMessageBox.critical(self, "错误", f"删除列失败：{str(e)}")

    @pyqtSlot()
    def on_load_data(self):
        if self.df is None:
            QMessageBox.warning(self, "警告", "请先加载数据")
            return

        self.ui.btn_Transfer.setEnabled(False)
        self.ui.progressBar.setValue(0)

        self.transfer_thread = DataTransferThread(self.node, self.df)
        self.transfer_thread.progress_update.connect(self.ui.progressBar.setValue)
        self.transfer_thread.transfer_done.connect(self.on_transfer_done)
        self.transfer_thread.start()

    @pyqtSlot(bool, str)
    def on_transfer_done(self, success: bool, msg: str):
        self.ui.btn_Transfer.setEnabled(True)
        self.ui.progressBar.setValue(100)
        if success:
            QMessageBox.information(self, "成功", msg)
        else:
            QMessageBox.critical(self, "错误", msg)

    def closeEvent(self, event):
        if self.loader_thread and self.loader_thread.isRunning():
            self.loader_thread.stop()
            self.loader_thread.wait()
        if self.transfer_thread and self.transfer_thread.isRunning():
            self.transfer_thread.wait()
        event.accept()


if __name__ == "__main__":
    import sys
    from PyQt5.QtWidgets import QApplication

    app = QApplication(sys.argv)

    class MockScene:
        def __init__(self):
            self.history = None

    mock_scene = MockScene()
    mock_node = ButtonNode(mock_scene, "Excel Import Node")
    window = Win_ExcelImport(mock_node)
    window.show()
    sys.exit(app.exec_())
